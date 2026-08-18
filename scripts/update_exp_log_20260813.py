"""
실험로그 xlsx 갱신 (2026-08-13).

세 가지를 반영한다.
  1. exp_003 — 학습이 08-12 14:29 에 2 epoch 로 끝났다. 진행 상황을 실제 값으로.
  2. exp_004 — 파트장 선별 8천건 학습 건을 '계획' 으로 추가.
  3. exp_001/002 — 해상도 이슈 검증 결과를 메모에 남겨 기준선 재실험을 막는다.

서식은 기존 행에서 그대로 복사한다 (상태별 행 색만 다르게).

usage:
    python3 scripts/update_exp_log_20260813.py
"""
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "scripts/LUXIA_VLM_학습현황.xlsx"
SHEET = "실험로그"
HEAD_ROW, DATA_ROW = 2, 3

# 상태별 행 색. 기존 시트가 쓰던 값 그대로 (완료 연두 / 진행 노랑).
STATE_FILL = {"완료": "E2EFDA", "진행": "FFF2CC", "계획": "DDEBF7", "폐기": "F2F2F2"}

RES_NOTE = (
    "\n[해상도 검증 08-13] json 해상도 ≠ 이미지 해상도 이슈 확인함. "
    "크롭 생성(build_eval_crops.py)이 축별 스케일 보정(sx=W/pw, sy=H/ph)을 "
    "이미 적용하고 있어 크롭 위치는 정합. val+test 99장 전량이 json≠이미지"
    "(배율 0.184~2.721, 비등방 2장)지만 보정 후 자름. 본 수치 재검증 불필요."
)

EXP003 = {
    "소요시간": "6시간 24분 (08-12 08:04:38~14:29 · train_runtime 23,060초)",
    "total_step": 1124,
    "메모": (
        "train_vlm.py 이미지 버그 수정 후 재학습.\n"
        "[08-13 확인] 학습 2 epoch 정상 종료(08-12 14:29). "
        "train_loss 0.0857 · 어댑터 ml/runs/upload-mix-qwen36 "
        "(checkpoint-1000 / checkpoint-1124 저장됨). "
        "GPU 2·3 은 현재 비어 있음 — 학습 끝난 상태.\n"
        "남은 일: 어댑터로 크롭 1,077건 평가 → CER/TEDS 채우면 '완료'. "
        "평가 전이라 상태는 '진행' 유지."
    ),
}

EXP004 = {
    "exp_id": "exp_004",
    "날짜": "2026-08-13",
    "상태": "계획",
    "가설/목적": (
        "파트장 선별 8천건(51만 중 추출)으로 학습하면 크롭 기준선"
        "(exp_001 LUXIA 0.4532 / exp_002 freeze 0.4221)을 넘는가"
    ),
    "모델": "Qwen3.6-35B-A3B\n(exp_003 과 같은 베이스)",
    "모델경로": "ml/models/Qwen3.6-35B-A3B",
    "어댑터": "예정: ml/runs/mix8k-qwen36",
    "train_set": (
        "김동영 파트장 선별 8천건\n"
        "kogovdoc 200 / multimodal 2,500~3,000 / pubtabnet 4,000~5,000"
    ),
    "n_train": "8,000 (수령 후 확정)",
    "n_val": 98,
    "학습방식": "LoRA r=16 (alpha 32, dropout 0.05) — exp_003 과 동일 조건",
    "freeze_vit": "TRUE",
    "freeze_aligner": "FALSE",
    "LoRA대상수": 250,
    "학습파라미터": "49.8M / 35.16B (0.142%)",
    "lr": 0.0001,
    "scheduler": "cosine (warmup 0.05)",
    "epochs": 2,
    "batch": 2,
    "grad_accum": 4,
    "유효배치": 8,
    "max_length": 8192,
    "max_px": 1000000,
    "grad_ckpt": "off",
    "group_by_len": "FALSE",
    "GPU": "H100 x2 (2,3) 예정",
    "평가셋": "영수증 크롭 1,077건 (text 973 / table 104) — exp_001~003 과 동일",
    "메모": (
        "[08-13 00:43 기준] train.jsonl 미수령 — 서버에서 못 찾음. "
        "받는 즉시 n_train·경로·total_step 확정 필요.\n"
        "원천은 vlm_data/Distill-VLM/.../01_taxonomy 아래 kogovdoc-bench, "
        "pubtabnet-html 로 보임.\n"
        "비교를 위해 학습 조건은 exp_003 과 같게 두고 데이터만 바꾼다."
    ),
}


def col_index(ws):
    return {ws.cell(HEAD_ROW, c).value: c
            for c in range(1, ws.max_column + 1) if ws.cell(HEAD_ROW, c).value}


def find_row(ws, exp_id):
    for r in range(DATA_ROW, ws.max_row + 1):
        if ws.cell(r, 1).value == exp_id:
            return r
    return None


def blank_row(ws):
    for r in range(DATA_ROW, ws.max_row + 1):
        if all(ws.cell(r, c).value in (None, "")
               for c in range(1, ws.max_column + 1)):
            return r
    return ws.max_row + 1


def style_row(ws, dst, src, state):
    """기존 행 서식을 그대로 옮기고 행 색만 상태에 맞춘다."""
    fill = PatternFill("solid", fgColor=STATE_FILL.get(state, "FFFFFF"))
    for c in range(1, ws.max_column + 1):
        s, d = ws.cell(src, c), ws.cell(dst, c)
        d.font, d.border = copy(s.font), copy(s.border)
        d.alignment, d.number_format = copy(s.alignment), s.number_format
        d.fill = fill
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def append_note(ws, idx, exp_id, text):
    r = find_row(ws, exp_id)
    if r is None:
        print(f"  ! {exp_id} 없음 — 건너뜀")
        return
    c = ws.cell(r, idx["메모"])
    cur = (c.value or "").rstrip()
    if "해상도 검증 08-13" in cur:
        print(f"  = {exp_id} 메모 이미 있음")
        return
    c.value = cur + text
    print(f"  + {exp_id} (행 {r}) 메모에 해상도 검증 결과 추가")


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    idx = col_index(ws)

    print("1) exp_003 학습 종료 반영")
    r3 = find_row(ws, "exp_003")
    for k, v in EXP003.items():
        ws.cell(r3, idx[k], v)
    print(f"  + exp_003 (행 {r3}) 소요시간·total_step·메모 갱신")

    print("2) exp_004 추가")
    r4 = find_row(ws, "exp_004") or blank_row(ws)
    style_row(ws, r4, r3, EXP004["상태"])
    for k, v in EXP004.items():
        if k not in idx:
            raise SystemExit(f"없는 열: {k}")
        ws.cell(r4, idx[k], v)
    print(f"  + exp_004 (행 {r4}) 계획으로 추가")

    print("3) 해상도 검증 결과 메모")
    for eid in ("exp_001", "exp_002"):
        append_note(ws, idx, eid, RES_NOTE)

    wb.save(XLSX)
    print(f"\n저장: {XLSX} ({ws.max_row}행 x {ws.max_column}열)")


if __name__ == "__main__":
    main()
