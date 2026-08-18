"""
실험로그 xlsx 에 크롭 평가 결과를 채워 넣는다.

평가가 끝날 때마다 다시 돌리면 해당 exp_id 행만 갱신된다 (없으면 추가).
학습이 끝나고 어댑터 평가를 돌린 뒤에도 같은 명령으로 채우면 된다.

원본 시트에 없는 지표(TEDS / AVG / CER 중앙값)는 '결과' 그룹 뒤에 열을 만들어
붙이고, 컬럼가이드 시트에도 설명을 같이 넣는다. 기존 열·서식·수식은 건드리지
않는다.

usage:
    # 평가 결과로 지표 채우기
    python3 scripts/fill_exp_log.py \
        --exp exp_003=receipt_data/review/eval_crops_luxia.json
    # 학습중 체크포인트처럼 JSON 안에 모델이 여럿일 때
    python3 scripts/fill_exp_log.py \
        --exp exp_004=receipt_data/review/eval_crops_upload_mix.json --model-key checkpoint-1000
    # 임의 열 채우기
    python3 scripts/fill_exp_log.py --set "exp_004:상태=진행" "exp_004:epochs=2"
"""
import argparse
import json
import statistics as st
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "scripts/LUXIA_VLM_학습현황.xlsx"
SHEET = "실험로그"
HEAD_ROW, DATA_ROW = 2, 3

# 요청 지표(CER / TEDS / AVG)를 담을 열. 원본 '결과' 그룹에 TEDS·AVG 가 없다.
NEW_COLS = [
    ("평가셋", "결과", "크롭 평가 대상. 예: 영수증 크롭 1,077건 (text 973 / table 104)"),
    ("평가건수", "결과", "평가에 실제로 쓴 크롭 수"),
    ("CER_중앙값", "결과", "텍스트 크롭 CER 의 중앙값. CER 은 위로 열린 값이라 폭주한 "
                       "소수 건이 평균을 지배한다 — 대표값은 중앙값을 쓴다"),
    ("CER_초과1건수", "결과", "CER>1 인 건수 (출력이 정답보다 길어진 폭주). 평균이 나쁠 때 "
                          "이 값을 같이 봐야 원인이 보인다"),
    ("TEDS", "결과", "표 크롭 TEDS 평균. 높을수록 좋음"),
    ("AVG_매크로", "결과", "텍스트 (1−CER) 과 표 TEDS 의 평균 (두 과제 평균)"),
    ("AVG_마이크로", "결과", "크롭 단위 평균. text 973 / table 104 로 9배 차이라 매크로와 같이 본다"),
    ("평가결과_json", "기타", "원본 결과 JSON 경로 (크롭별 출력까지 들어 있음)"),
]


def metrics(path, model_key=None):
    d = json.loads(Path(path).read_text(encoding="utf-8"))
    names = list(d["summary"])
    name = model_key or names[-1]
    if name not in d["summary"]:
        raise SystemExit(f"{path} 안에 '{name}' 없음. 있는 것: {names}")
    s = d["summary"][name]
    cs = [r[name]["cer"] for r in d["rows"]
          if r["task"] == "text" and r.get(name, {}).get("cer") is not None]
    txt, tbl, avg = s.get("text", {}), s.get("table", {}), s.get("AVG", {})
    out = {
        "평가셋": f'영수증 크롭 {d["n"]}건 '
                f'(text {txt.get("n","-")} / table {tbl.get("n","-")})',
        "평가건수": d["n"],
        "평가결과_json": str(Path(path)),
    }
    if "CER" in txt:
        out["CER_전체"] = round(txt["CER"], 4)
    if cs:
        out["CER_중앙값"] = round(st.median(cs), 4)
        out["CER_초과1건수"] = sum(1 for c in cs if c > 1)
    if "TEDS" in tbl:
        out["TEDS"] = round(tbl["TEDS"], 4)
    if "매크로" in avg:
        out["AVG_매크로"] = round(avg["매크로"], 4)
    if "마이크로" in avg:
        out["AVG_마이크로"] = round(avg["마이크로"], 4)
    return name, out


def ensure_cols(ws):
    """없는 열만 뒤에 붙이고 이름 -> 열번호 맵을 낸다."""
    idx = {ws.cell(HEAD_ROW, c).value: c
           for c in range(1, ws.max_column + 1) if ws.cell(HEAD_ROW, c).value}
    src = ws.cell(HEAD_ROW, 1)
    for name, group, _ in NEW_COLS:
        if name in idx:
            continue
        c = ws.max_column + 1
        ws.cell(1, c, group).font = Font(bold=True)
        h = ws.cell(HEAD_ROW, c, name)
        h.font = Font(bold=True) if src.font is None else Font(bold=True)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[h.column_letter].width = max(11, min(34, len(name) + 8))
        idx[name] = c
    return idx


def guide(wb):
    """컬럼가이드 시트에도 새 열 설명을 추가한다 (없는 것만)."""
    if "컬럼가이드" not in wb.sheetnames:
        return
    ws = wb["컬럼가이드"]
    have = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
    r = ws.max_row + 1
    for name, _, desc in NEW_COLS:
        if name in have:
            continue
        ws.cell(r, 1, name)
        c = ws.cell(r, 2, desc)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1


def row_for(ws, exp_id):
    for r in range(DATA_ROW, ws.max_row + 1):
        if ws.cell(r, 1).value == exp_id:
            return r
    for r in range(DATA_ROW, ws.max_row + 1):          # 빈 줄 재사용
        if all(ws.cell(r, c).value in (None, "")
               for c in range(1, ws.max_column + 1)):
            ws.cell(r, 1, exp_id)
            return r
    r = ws.max_row + 1
    ws.cell(r, 1, exp_id)
    return r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--exp", nargs="*", default=[],
                    help="exp_id=평가JSON 형식. 여러 개 가능")
    ap.add_argument("--model-key", default=None,
                    help="JSON 안에 모델이 여럿일 때 고를 이름 (예: checkpoint-1000)")
    ap.add_argument("--set", nargs="*", default=[],
                    help="exp_id:열이름=값")
    ap.add_argument("--xlsx", default=str(XLSX))
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx)
    ws = wb[SHEET]
    idx = ensure_cols(ws)
    guide(wb)

    for spec in args.exp:
        exp_id, _, path = spec.partition("=")
        name, m = metrics(path, args.model_key)
        r = row_for(ws, exp_id)
        for k, v in m.items():
            if k not in idx:
                raise SystemExit(f"없는 열: {k}")
            ws.cell(r, idx[k], v)
        print(f"{exp_id} (행 {r}): {name}  CER평균 {m.get('CER_전체')} "
              f"중앙값 {m.get('CER_중앙값')} TEDS {m.get('TEDS')} "
              f"AVG {m.get('AVG_매크로')}")

    for spec in args.set:
        eid, _, rest = spec.partition(":")
        col, _, val = rest.partition("=")
        if col not in idx:
            raise SystemExit(f"없는 열: {col}\n있는 열: {list(idx)}")
        try:                                   # 숫자로 넣을 수 있으면 숫자로
            val = int(val) if val.lstrip("-").isdigit() else float(val)
        except ValueError:
            pass
        ws.cell(row_for(ws, eid), idx[col], val)

    wb.save(args.xlsx)
    print(f"저장: {args.xlsx} ({ws.max_row}행 x {ws.max_column}열)")


if __name__ == "__main__":
    main()
