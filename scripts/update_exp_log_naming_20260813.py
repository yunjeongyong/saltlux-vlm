"""
실험로그 xlsx — 명명 규칙 적용 + 학습 로그 점검 결과 반영 (2026-08-13).

상무님 지적 2건 대응.
  ① 모델 호칭을 base / serve / ft 3역할로 정규화. "freeze 모델" 폐기.
  ② exp_003 의 과적합(eval_loss 단조 증가)을 로그에 남긴다.

확정된 비교 실험 스펙(평가 200건 = 영수증 100 + 계산서 100, 지표 CER)도
exp_004 에 반영한다.

usage:
    python3 scripts/update_exp_log_naming_20260813.py
"""
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "scripts/LUXIA_VLM_학습현황.xlsx"
SHEET = "실험로그"
HEAD_ROW, DATA_ROW = 2, 3

# ① 모델 호칭 정규화. {계열}-{크기}-{역할}[-{데이터}][-{버전}]
#    base = 원본 가중치 / serve = 사내 서빙 설정 / ft = 우리 어댑터
NAMES = {
    "exp_001": "qwen36-a3b-serve-luxiadp-high\n"
               "(구 호칭: 베이스 모델 · LUXIA)\n"
               "※ base 와 동일 가중치라는 전제는 미검증",
    "exp_002": "qwen36-a3b-base\n"
               "(구 호칭: freeze 모델)\n"
               "원본 가중치 · 어댑터 없음 · 추론만",
    "exp_003": "qwen36-a3b-ft-uploadmix-v1\n"
               "(구 호칭: freeze 모델에 LoRA)\n"
               "base + ml/runs/upload-mix-qwen36",
    "exp_004": "qwen36-a3b-ft-mix8k-v1\n"
               "base + (어댑터 예정)",
}

# ② exp_003 과적합. eval_loss 는 4회 측정 전부를 남겨야 판단이 된다.
EXP003 = {
    "eval_loss": "0.1628 → 0.1851 → 0.2163 → 0.2294\n"
                 "(step 250/500/750/1000 · 단조 증가 = 과적합)",
    "메모": (
        "train_vlm.py 이미지 버그 수정 후 재학습.\n"
        "[08-13 확인] 학습 2 epoch 정상 종료(08-12 14:29). train_loss 0.0857 · "
        "어댑터 ml/runs/upload-mix-qwen36. GPU 2·3 현재 비어 있음.\n"
        "[학습로그 점검] ⚠ 과적합 확인 — train loss 0.164→0.054(-67%) 인데 "
        "eval_loss 0.163→0.229(+41%) 로 단조 증가. 최저는 step 250(epoch 0.44).\n"
        "⚠ load_best_model_at_end=False + save_steps=500 + save_total_limit=2 라 "
        "best(step 250) 는 저장 안 됐고 남은 건 ckpt-1000/1124 뿐 — 최종 어댑터가 "
        "최악 구간이다. 평가는 ckpt-1000 과 최종을 둘 다 재야 한다.\n"
        "원인 의심: 영수증 187장×8회 반복 → 2 epoch 면 같은 장을 16번 봄.\n"
        "정상 확인: NaN/Inf/OOM/토큰초과 0건 · grad_norm 0.096~0.181 안정 · "
        "LoRA 250개(linear_attn 포함) 정상 · 누수 0.\n"
        "미확인: 로그에 'vision encoder 동결 파라미터 0개' — 이름만 동결인지 "
        "이미 동결이라 잡을 게 없었는지 train_vlm.py 확인 필요.\n"
        "상세: 노션_모델명명규칙_및_학습로그점검_20260813.md"
    ),
}

# 확정된 비교 실험 스펙을 exp_004 에 반영
EXP004 = {
    "평가셋": "영수증 100 + 계산서 100 = 200건 (확정 스펙)\n"
            "※ 크롭 1,077건(exp_001~003)과 다른 평가셋 — 직접 비교 불가",
    "epochs": 1,
    "메모": (
        "[08-13 00:43 기준] train.jsonl 미수령 — 서버에서 못 찾음. 받는 즉시 "
        "n_train·경로·total_step 확정 필요.\n"
        "원천은 vlm_data/Distill-VLM/.../01_taxonomy 아래 kogovdoc-bench, "
        "pubtabnet-html 로 보임.\n"
        "[exp_003 과적합 반영] 같은 설정으로 돌리면 같은 과적합이 난다. 바꿀 것 — "
        "①epochs 2→1 ②load_best_model_at_end=True ③save_steps=eval_steps=250 "
        "④영수증 반복 8회→2~3회.\n"
        "평가 지표는 CER (확정). 평가셋이 크롭 1,077건과 달라 exp_001~003 수치와 "
        "나란히 놓으면 안 된다 — 비교하려면 base/serve 도 200건으로 다시 재야 함."
    ),
}

# 컬럼가이드에 명명 규칙을 박아둔다. 다음에 또 "freeze 모델이 뭐냐" 가 안 나오게.
GUIDE = [
    ("모델 [명명규칙]",
     "{계열}-{크기}-{역할}[-{데이터}][-{버전}]. 역할 3가지만 쓴다 — "
     "base(원본 가중치, 어댑터 없음) / serve(같은 가중치 + 사내 서빙 설정) / "
     "ft(base + 우리 LoRA 어댑터). "
     '"freeze 모델" 은 폐기 — freeze 는 freeze_vit/freeze_aligner 열에서 '
     "'파라미터 동결' 뜻으로만 쓴다 (2026-08-13, 상무님 지적 반영)"),
    ("eval_loss",
     "검증셋 loss. 1회만 적지 말고 측정 전부를 적는다 — 값 하나로는 "
     "과적합(train↓ eval↑)이 안 보인다. exp_003 이 그 사례"),
]


def col_index(ws):
    return {ws.cell(HEAD_ROW, c).value: c
            for c in range(1, ws.max_column + 1) if ws.cell(HEAD_ROW, c).value}


def find_row(ws, exp_id):
    for r in range(DATA_ROW, ws.max_row + 1):
        if ws.cell(r, 1).value == exp_id:
            return r
    return None


def apply(ws, idx, exp_id, fields):
    r = find_row(ws, exp_id)
    if r is None:
        print(f"  ! {exp_id} 없음 — 건너뜀")
        return
    for k, v in fields.items():
        if k not in idx:
            raise SystemExit(f"없는 열: {k}")
        ws.cell(r, idx[k], v)
    print(f"  + {exp_id} (행 {r}) {', '.join(fields)} 갱신")


def guide(wb):
    if "컬럼가이드" not in wb.sheetnames:
        return
    ws = wb["컬럼가이드"]
    have = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1)}
    r = ws.max_row + 1
    for name, desc in GUIDE:
        if name in have:
            print(f"  = 컬럼가이드 '{name}' 이미 있음")
            continue
        ws.cell(r, 1, name)
        c = ws.cell(r, 2, desc)
        c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
        print(f"  + 컬럼가이드 '{name}' 추가")
        r += 1


def main():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb[SHEET]
    idx = col_index(ws)

    print("① 모델 호칭 정규화")
    for eid, name in NAMES.items():
        apply(ws, idx, eid, {"모델": name})

    print("② 학습 로그 점검 결과")
    apply(ws, idx, "exp_003", EXP003)

    print("③ 확정 스펙 반영 (평가 200건 / CER)")
    apply(ws, idx, "exp_004", EXP004)

    print("④ 컬럼가이드")
    guide(wb)

    wb.save(XLSX)
    print(f"\n저장: {XLSX}")


if __name__ == "__main__":
    main()
