"""LUXIA_VLM_04 프롬프트 정리 — 01~03 과 같은 서식·색 팔레트를 쓴다."""
import json, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TH=Side(style="thin",color="BFBFBF"); BD=Border(left=TH,right=TH,top=TH,bottom=TH)
HF=PatternFill("solid",fgColor="D9E2F3")   # 소제목 헤더
GF=PatternFill("solid",fgColor="B4C7E7")   # 그룹 헤더 (병합)
NEW=PatternFill("solid",fgColor="FFF2CC")  # 노랑 — 실패 억제 지시문
HI=PatternFill("solid",fgColor="E2EFDA")   # 초록 — 배포 문구
BLU=PatternFill("solid",fgColor="DDEBF7")  # 파랑 — 문구 변형
C=Alignment(horizontal="center",vertical="center",wrap_text=True)
L=Alignment(horizontal="left",vertical="top",wrap_text=True)
LC=Alignment(horizontal="left",vertical="center",wrap_text=True)
ROOT="/data/workspace/yyj/data"

def W(ws,w):
    for i,x in enumerate(w,1): ws.column_dimensions[get_column_letter(i)].width=x
def h(ws,r,c,v,fill=HF):
    x=ws.cell(r,c,v); x.font=Font(bold=True,size=10); x.fill=fill; x.border=BD; x.alignment=C
def cl(ws,r,c,v,fill=None,align=None,bold=False,fmt=None):
    x=ws.cell(r,c,v); x.border=BD; x.alignment=align or C
    if fill: x.fill=fill
    if bold: x.font=Font(bold=True,size=10)
    if fmt: x.number_format=fmt
def note(ws,r,c2,txt):
    cl(ws,r,2,txt,align=LC); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=c2)
def legend(ws,r,items):
    x=ws.cell(r,2,"범례"); x.font=Font(bold=True,size=9); x.alignment=C
    c=3
    for fill,label in items:
        a=ws.cell(r,c,""); a.fill=fill; a.border=BD
        b=ws.cell(r,c+1,label); b.alignment=LC; b.font=Font(size=9)
        c+=2
    return r+1

P=json.load(open("/tmp/prompts.json"))
TOT=sum(o["n"] for o in P)
GROUP={0:"배포 문구",1:"실패 억제 지시문",2:"배포 문구",3:"배포 문구",4:"실패 억제 지시문"}
FILLOF={"배포 문구":HI,"실패 억제 지시문":NEW,"문구 변형":BLU}

wb=Workbook()

# ══ 1. 프롬프트 인벤토리 ══
ws=wb.active; ws.title="1.프롬프트 인벤토리"
W(ws,[3,5,16,76,28,30,9,8,6,8])
cl(ws,2,2,f"exp_012 학습셋 프롬프트 — 고유 지시문 {len(P)}종 / {TOT:,}행  (2026-09-02)",align=LC,bold=True)
ws.merge_cells("B2:J2")
h(ws,3,2,"#",GF); ws.merge_cells("B3:B4")
h(ws,3,3,"분류",GF); ws.merge_cells("C3:C4")
h(ws,3,4,"프롬프트 지시문 (전문)",GF); ws.merge_cells("D3:D4")
h(ws,3,5,"사용 태스크 (문서 유형)",GF); ws.merge_cells("E3:E4")
h(ws,3,6,"주요 출처",GF); ws.merge_cells("F3:F4")
h(ws,3,7,"규모",GF); ws.merge_cells("G3:H3")
h(ws,4,7,"행 수"); h(ws,4,8,"비중")
h(ws,3,9,"지시문 길이",GF); ws.merge_cells("I3:J3")
h(ws,4,9,"줄"); h(ws,4,10,"글자")
ws.freeze_panes="D5"
r=5
for i,o in enumerate(P):
    grp=GROUP.get(i,"문구 변형"); f=FILLOF[grp]
    cl(ws,r,2,i+1,f); cl(ws,r,3,grp,f)
    cl(ws,r,4,o["prompt"],f,align=L)
    cl(ws,r,5," · ".join(f"{k} ({v:,})" for k,v in o["tasks"]),f,align=L)
    cl(ws,r,6," · ".join(k for k,_ in o["src"][:3]),f,align=L)
    cl(ws,r,7,o["n"],f,fmt="#,##0"); cl(ws,r,8,o["pct"]/100,f,fmt="0.00%")
    cl(ws,r,9,o["lines"],f); cl(ws,r,10,o["chars"],f,fmt="#,##0")
    ws.row_dimensions[r].height=min(320,16+o["prompt"].count("\n")*11+len(o["prompt"])//76*11)
    r+=1
cl(ws,r,4,"합계",bold=True,align=LC)
cl(ws,r,7,TOT,bold=True,fmt="#,##0"); cl(ws,r,8,1,bold=True,fmt="0.00%")
for c in (2,3,5,6,9,10): cl(ws,r,c,"")
r+=2
r=legend(ws,r,[(HI,"배포 문구 — DocStudio 가 실제로 보내는 문구"),
               (NEW,"실패 억제 지시문 — 반복 폭주·표 붕괴 방지"),
               (BLU,"문구 변형 — 같은 뜻 다른 표현")])
r+=1
note(ws,r,10,"※ 프롬프트는 별도 설정 파일이 아니라 train.jsonl 의 messages[0].content 에 들어 있다. "
             "train_vlm.py 는 그 값을 챗 템플릿에 끼울 뿐이다 (113~133행, enable_thinking=False). "
             "따라서 '이 모델이 무슨 프롬프트로 학습됐나' 의 확정 위치는 train.jsonl 이다.")

# ══ 2. 태스크별 요약 ══
ws=wb.create_sheet("2.태스크별 요약")
W(ws,[3,34,11,11,9,60])
cl(ws,2,2,"태스크(문서 유형)별 프롬프트 사용 현황",align=LC,bold=True); ws.merge_cells("B2:F2")
h(ws,3,2,"태스크 (문서 유형)",GF); h(ws,3,3,"프롬프트 종수",GF)
h(ws,3,4,"행 수",GF); h(ws,3,5,"비중",GF); h(ws,3,6,"대표 지시문",GF)
ws.freeze_panes="C4"
byt=collections.defaultdict(lambda:{"n":0,"p":set(),"top":("",0)})
for o in P:
    for t,c in o["tasks"]:
        b=byt[t]; b["n"]+=c; b["p"].add(o["prompt"])
        if c>b["top"][1]: b["top"]=(o["prompt"],c)
r=4
for t,b in sorted(byt.items(), key=lambda x:-x[1]["n"]):
    f=BLU if len(b["p"])>1 else None
    cl(ws,r,2,t,f,align=LC); cl(ws,r,3,len(b["p"]),f)
    cl(ws,r,4,b["n"],f,fmt="#,##0"); cl(ws,r,5,b["n"]/TOT,f,fmt="0.00%")
    s=b["top"][0].replace("\n"," ")
    cl(ws,r,6,s[:108]+("…" if len(s)>108 else ""),f,align=LC)
    r+=1
cl(ws,r,2,"합계",bold=True,align=LC); cl(ws,r,3,len(P),bold=True)
cl(ws,r,4,TOT,bold=True,fmt="#,##0"); cl(ws,r,5,1,bold=True,fmt="0.00%"); cl(ws,r,6,"")
r+=2
r=legend(ws,r,[(BLU,"프롬프트 2종 이상 — 문구 변형 적용")]); r+=1
note(ws,r,6,"※ doc_parsing/* 12종은 전부 같은 프롬프트 1종을 쓴다. Document Studio 는 문서 종류를 모르고 보내므로 "
            "도메인별로 다른 문구를 주면 배포 때 쓸 수 없다.")

# ══ 3. 분류 요약 ══
ws=wb.create_sheet("3.분류 요약")
W(ws,[3,22,10,11,9,58])
cl(ws,2,2,"프롬프트 분류 — 세 갈래",align=LC,bold=True); ws.merge_cells("B2:F2")
for i,n in enumerate(["분류","종수","행 수","비중","의도"],2): h(ws,3,i,n,GF)
g=collections.Counter(); gs=collections.Counter()
for i,o in enumerate(P):
    k=GROUP.get(i,"문구 변형"); g[k]+=o["n"]; gs[k]+=1
INTENT={"배포 문구":"Document Studio 가 실제로 보내는 문구를 그대로 학습한다. 신규 도메인 12종이 전부 이 중 하나를 쓴다.",
        "실패 억제 지시문":"반복 폭주와 표 붕괴를 막으려 규칙을 나열했다. 20줄·59줄로 길다.",
        "문구 변형":"같은 뜻을 다르게 써 한 문구에만 반응하지 않게 한다."}
r=4
for k in ["배포 문구","실패 억제 지시문","문구 변형"]:
    cl(ws,r,2,k,FILLOF[k],align=LC,bold=True); cl(ws,r,3,gs[k],FILLOF[k])
    cl(ws,r,4,g[k],FILLOF[k],fmt="#,##0"); cl(ws,r,5,g[k]/TOT,FILLOF[k],fmt="0.0%")
    cl(ws,r,6,INTENT[k],FILLOF[k],align=LC); r+=1
cl(ws,r,2,"합계",bold=True,align=LC); cl(ws,r,3,len(P),bold=True)
cl(ws,r,4,TOT,bold=True,fmt="#,##0"); cl(ws,r,5,1,bold=True,fmt="0.0%"); cl(ws,r,6,"")
r+=2
cl(ws,r,2,"프롬프트가 정의된 위치",align=LC,bold=True); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=6); r+=1
for i,n in enumerate(["파일","담고 있는 것","성격"],2): h(ws,r,i,n,GF)
ws.merge_cells(start_row=r,start_column=4,end_row=r,end_column=6); ws.cell(r,5).fill=GF; ws.cell(r,6).fill=GF
ws.cell(r,5).border=BD; ws.cell(r,6).border=BD
r+=1
FILES=[("scripts/prompts_receipt.py","OCR_PROMPT · TABLE_PROMPT (긴 지시문 2종)","모듈 · PROMPT_VERSION = 2026-08-12"),
       ("scripts/prompt_templates.py","FORMAT_RULE · SCOPE_RULE · TASKS 조합 규칙","모듈"),
       ("scripts/gt_prompts.py","PAGE_KIE_SCHEMA · 정답 생성용 프롬프트","모듈 (학습 아님)"),
       ("ml/prep_exp012_restart.sh 외 4개","짧은 배포 문구 3종","⚠ 하드코딩 · 같은 문구가 여러 파일에 복사됨")]
for a,b,c in FILES:
    cl(ws,r,2,a,align=LC); cl(ws,r,3,b,align=LC)
    ws.merge_cells(start_row=r,start_column=3,end_row=r,end_column=4); ws.cell(r,4).border=BD
    cl(ws,r,5,c,align=LC); ws.merge_cells(start_row=r,start_column=5,end_row=r,end_column=6); ws.cell(r,6).border=BD
    r+=1
r+=1
note(ws,r,6,"※ 짧은 배포 문구 3종이 모듈이 아니라 4~5개 스크립트에 각각 하드코딩되어 있다. "
            "하나 고치려면 전부 찾아야 하므로 prompts_receipt.py 한 곳으로 모으는 것을 제안한다.")

p=f"{ROOT}/LUXIA_VLM_04_프롬프트정리.xlsx"
wb.save(p); print("저장:",p)
for n in wb.sheetnames:
    w=wb[n]; rows=sum(1 for rr in range(1,w.max_row+1) if any(w.cell(rr,c).value is not None for c in range(1,w.max_column+1)))
    print(f"  [{n}]  {rows}행 × {w.max_column}열")
